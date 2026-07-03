#!/usr/bin/env python3
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/zileitian/Desktop/Phenotyping Capsaicin-Induced Visceral Pain Dynamics")
MANUSCRIPT = ROOT / "MANUSCRIPT.md"
SUPPLEMENTARY = ROOT / "supplementary_latest.md"
OUTPUT = ROOT / "Temporal_Dynamics_Capsaicin_latest_manuscript_supplementary_2026-07-03.docx"
TABLE_GEOMETRY = Path(
    "/Users/zileitian/.codex/plugins/cache/openai-primary-runtime/documents/26.630.12135/skills/documents/scripts/table_geometry.py"
)


def load_table_geometry():
    import importlib.util

    spec = importlib.util.spec_from_file_location("table_geometry", TABLE_GEOMETRY)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


table_geometry = load_table_geometry()


def pandoc_ast(path: Path) -> dict:
    raw = subprocess.check_output(
        ["pandoc", str(path), "--from", "markdown+pipe_tables", "--to", "json"],
        text=True,
    )
    return json.loads(raw)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    npf = normal.paragraph_format
    npf.space_before = Pt(0)
    npf.space_after = Pt(8)
    npf.line_spacing = 1.15
    npf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    h1 = doc.styles["Heading 1"]
    h1.font.size = Pt(20)
    h1.font.bold = False
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1pf = h1.paragraph_format
    h1pf.space_before = Pt(20)
    h1pf.space_after = Pt(6)

    h2 = doc.styles["Heading 2"]
    h2.font.size = Pt(16)
    h2.font.bold = False
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2pf = h2.paragraph_format
    h2pf.space_before = Pt(18)
    h2pf.space_after = Pt(6)

    h3 = doc.styles["Heading 3"]
    h3.font.size = Pt(14)
    h3.font.bold = False
    h3.font.color.rgb = RGBColor(67, 67, 67)
    h3pf = h3.paragraph_format
    h3pf.space_before = Pt(16)
    h3pf.space_after = Pt(4)


def stringify_inlines(inlines: Iterable[dict]) -> str:
    parts: list[str] = []
    for item in inlines:
        t = item["t"]
        c = item.get("c")
        if t == "Str":
            parts.append(c)
        elif t == "Space":
            parts.append(" ")
        elif t in {"SoftBreak", "LineBreak"}:
            parts.append("\n")
        elif t in {"Emph", "Strong", "Strikeout", "SmallCaps", "Underline"}:
            parts.append(stringify_inlines(c))
        elif t == "Code":
            parts.append(c[1])
        elif t == "Link":
            parts.append(stringify_inlines(c[1]))
        elif t == "Quoted":
            parts.append(stringify_inlines(c[1]))
        elif t == "Math":
            parts.append(c[1])
        elif t == "Image":
            parts.append(stringify_inlines(c[1]))
    return "".join(parts)


def add_inlines(paragraph, inlines: Iterable[dict]) -> None:
    for item in inlines:
        t = item["t"]
        c = item.get("c")
        if t == "Str":
            paragraph.add_run(c)
        elif t == "Space":
            paragraph.add_run(" ")
        elif t in {"SoftBreak", "LineBreak"}:
            paragraph.add_run().add_break()
        elif t == "Strong":
            run = paragraph.add_run(stringify_inlines(c))
            run.bold = True
        elif t == "Emph":
            run = paragraph.add_run(stringify_inlines(c))
            run.italic = True
        elif t == "Underline":
            run = paragraph.add_run(stringify_inlines(c))
            run.underline = True
        elif t == "Code":
            run = paragraph.add_run(c[1])
            run.font.name = "Courier New"
        elif t == "Link":
            paragraph.add_run(stringify_inlines(c[1]))
        elif t == "Quoted":
            paragraph.add_run(stringify_inlines(c[1]))
        elif t == "Math":
            paragraph.add_run(c[1])
        elif t == "Image":
            # Images are handled at the block level.
            paragraph.add_run(stringify_inlines(c[1]))


def looks_like_image_para(block: dict) -> bool:
    if block["t"] not in {"Para", "Plain"}:
        return False
    inlines = [item for item in block["c"] if item["t"] != "Space"]
    return len(inlines) == 1 and inlines[0]["t"] == "Image"


def add_image_block(doc: Document, block: dict) -> None:
    image = [item for item in block["c"] if item["t"] == "Image"][0]
    target = image["c"][2][0]
    path = (ROOT / target).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.0))
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def add_figure_block(doc: Document, block: dict) -> None:
    body = block["c"][2]
    for item in body:
        if item["t"] not in {"Plain", "Para"}:
            continue
        for inline in item["c"]:
            if inline["t"] != "Image":
                continue
            target = inline["c"][2][0]
            path = (ROOT / target).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(path), width=Inches(6.0))
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            return


def add_para(doc: Document, block: dict, style: str = "Normal") -> None:
    p = doc.add_paragraph(style=style)
    add_inlines(p, block["c"])


def add_bullet_list(doc: Document, block: dict) -> None:
    for item in block["c"]:
        paras = item if isinstance(item, list) else []
        first = True
        for sub in paras:
            p = doc.add_paragraph(style="List Bullet")
            if not first:
                p.paragraph_format.left_indent = Inches(0.25)
            if sub["t"] in {"Para", "Plain"}:
                add_inlines(p, sub["c"])
            else:
                p.add_run(stringify_inlines(sub.get("c", [])))
            first = False


def header_text(block: dict) -> str:
    return stringify_inlines(block["c"][2])


def table_widths(headers: list[str]) -> list[int]:
    mapping = {
        ("Characteristic", "Value"): [5760, 3600],
        ("Symptom", "Region", "Weight"): [3600, 3960, 1800],
        ("Node", "Type", "Degree", "Weighted Degree"): [3240, 1800, 1800, 2520],
        (
            "Model",
            "Accuracy (mean +/- SD)",
            "Balanced Accuracy",
            "Macro F1",
            "Weighted F1",
        ): [2160, 2160, 1680, 1560, 1800],
        ("Safety outcome", "Number of participants", "Percentage"): [5400, 2160, 1800],
    }
    return mapping.get(tuple(headers), table_geometry.column_widths_from_weights([1] * len(headers)))


def cell_blocks_to_text(cell_blocks: list[dict]) -> str:
    return "\n".join(
        stringify_inlines(block["c"]) for block in cell_blocks if block["t"] in {"Para", "Plain"}
    ).strip()


def add_table(doc: Document, block: dict) -> None:
    _, _, _, head, bodies, _ = block["c"]
    head_rows = head[1]
    if not head_rows:
        return

    header_cells = head_rows[0][1]
    headers = [cell_blocks_to_text(cell[4]) for cell in header_cells]

    body_rows: list[list[str]] = []
    for body in bodies:
        for row in body[3]:
            texts = [cell_blocks_to_text(cell[4]) for cell in row[1]]
            body_rows.append(texts)

    table = doc.add_table(rows=1 + len(body_rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    widths = table_widths(headers)

    for col, text in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = text
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.style = doc.styles["Normal"]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.runs[0].bold = True

    for row_idx, row in enumerate(body_rows, start=1):
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.style = doc.styles["Normal"]
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)

    table_geometry.apply_table_geometry(
        table,
        widths,
        table_width_dxa=9360,
        indent_dxa=120,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120},
    )

    for col_idx, width in enumerate(widths):
        table.columns[col_idx].width = width

    for row in table.rows:
        for col_idx, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                if col_idx == len(widths) - 1 and paragraph.text.replace(".", "", 1).replace("%", "").isdigit():
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph()


def render_blocks(doc: Document, blocks: list[dict]) -> None:
    for block in blocks:
        t = block["t"]
        if t == "Header":
            level = block["c"][0]
            text = header_text(block)
            doc.add_heading(text, level=min(level, 3))
        elif t in {"Para", "Plain"} and looks_like_image_para(block):
            add_image_block(doc, block)
        elif t == "Figure":
            add_figure_block(doc, block)
        elif t in {"Para", "Plain"}:
            add_para(doc, block)
        elif t == "BulletList":
            add_bullet_list(doc, block)
        elif t == "Table":
            add_table(doc, block)
        elif t == "HorizontalRule":
            doc.add_paragraph()
        elif t == "CodeBlock":
            p = doc.add_paragraph()
            run = p.add_run(block["c"][1])
            run.font.name = "Courier New"
        elif t == "OrderedList":
            for item in block["c"][1]:
                for sub in item:
                    p = doc.add_paragraph(style="List Number")
                    if sub["t"] in {"Para", "Plain"}:
                        add_inlines(p, sub["c"])


def main() -> int:
    doc = Document()
    configure_document(doc)

    render_blocks(doc, pandoc_ast(MANUSCRIPT)["blocks"])
    doc.add_page_break()
    render_blocks(doc, pandoc_ast(SUPPLEMENTARY)["blocks"])

    doc.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
